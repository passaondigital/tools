"""
DOCX-Exporter für E-Books.

Dieses Modul verwendet python-docx zur Erstellung professioneller
Word-Dokumente aus E-Book-Daten.
"""

import re
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from ebook_generator.models.ebook import EBook, Seitenformat


class DOCXExporter:
    """
    Exportiert E-Books in das DOCX-Format.

    Diese Klasse verwendet python-docx, um professionelle Word-Dokumente
    mit Titelseite, Inhaltsverzeichnis und formatiertem Text zu erstellen.
    """

    # Seitenformate-Mapping (Breite, Höhe in cm)
    SEITENFORMATE = {
        Seitenformat.A4: (21.0, 29.7),
        Seitenformat.A5: (14.8, 21.0),
        Seitenformat.LETTER: (21.59, 27.94),
        Seitenformat.TASCHENBUCH: (12.7, 20.32),
        Seitenformat.GROSSDRUCK: (17.0, 24.0),
    }

    def __init__(self):
        """Initialisiert den DOCX-Exporter."""
        self.document: Optional[Document] = None

    def exportieren(self, ebook: EBook, ausgabepfad: Path) -> Path:
        """
        Exportiert das E-Book als DOCX.

        Args:
            ebook: Das zu exportierende E-Book
            ausgabepfad: Pfad für die Ausgabedatei

        Returns:
            Pfad zur erstellten DOCX-Datei
        """
        self.document = Document()

        # Stile einrichten
        self._stile_erstellen(ebook)

        # Seitenformat einrichten
        self._seitenformat_einrichten(ebook)

        # Dokumenteigenschaften setzen
        self._eigenschaften_setzen(ebook)

        # Titelseite
        if ebook.titelseite:
            self._titelseite_erstellen(ebook)
            self._seitenumbruch()

        # Copyright-Seite
        self._copyright_seite_erstellen(ebook)
        self._seitenumbruch()

        # Widmung
        if ebook.widmung:
            self._widmung_erstellen(ebook)
            self._seitenumbruch()

        # Inhaltsverzeichnis
        if ebook.inhaltsverzeichnis:
            self._inhaltsverzeichnis_erstellen(ebook)
            self._seitenumbruch()

        # Vorwort
        if ebook.vorwort:
            self._vorwort_erstellen(ebook)
            self._seitenumbruch()

        # Kapitel
        for i, kapitel in enumerate(ebook.kapitel):
            self._kapitel_erstellen(kapitel)
            if i < len(ebook.kapitel) - 1:
                self._seitenumbruch()

        # Nachwort
        if ebook.nachwort:
            self._seitenumbruch()
            self._nachwort_erstellen(ebook)

        # Dokument speichern
        self.document.save(str(ausgabepfad))

        return ausgabepfad

    def _stile_erstellen(self, ebook: EBook) -> None:
        """Erstellt benutzerdefinierte Stile."""
        styles = self.document.styles

        # Buch-Titel-Stil
        try:
            titel_stil = styles.add_style("BuchTitel", WD_STYLE_TYPE.PARAGRAPH)
        except ValueError:
            titel_stil = styles["BuchTitel"]
        titel_stil.font.size = Pt(28)
        titel_stil.font.bold = True
        titel_stil.font.color.rgb = RGBColor(26, 26, 26)
        titel_stil.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        titel_stil.paragraph_format.space_after = Pt(30)

        # Buch-Untertitel-Stil
        try:
            untertitel_stil = styles.add_style(
                "BuchUntertitel", WD_STYLE_TYPE.PARAGRAPH
            )
        except ValueError:
            untertitel_stil = styles["BuchUntertitel"]
        untertitel_stil.font.size = Pt(16)
        untertitel_stil.font.italic = True
        untertitel_stil.font.color.rgb = RGBColor(74, 74, 74)
        untertitel_stil.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        untertitel_stil.paragraph_format.space_after = Pt(20)

        # Autor-Stil
        try:
            autor_stil = styles.add_style("BuchAutor", WD_STYLE_TYPE.PARAGRAPH)
        except ValueError:
            autor_stil = styles["BuchAutor"]
        autor_stil.font.size = Pt(14)
        autor_stil.font.color.rgb = RGBColor(42, 42, 42)
        autor_stil.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        autor_stil.paragraph_format.space_after = Pt(50)

        # Kapitel-Titel-Stil
        try:
            kapitel_stil = styles.add_style(
                "KapitelTitel", WD_STYLE_TYPE.PARAGRAPH
            )
        except ValueError:
            kapitel_stil = styles["KapitelTitel"]
        kapitel_stil.font.size = Pt(20)
        kapitel_stil.font.bold = True
        kapitel_stil.font.color.rgb = RGBColor(26, 26, 26)
        kapitel_stil.paragraph_format.space_before = Pt(30)
        kapitel_stil.paragraph_format.space_after = Pt(20)
        kapitel_stil.paragraph_format.keep_with_next = True

        # Abschnitt-Titel-Stil
        try:
            abschnitt_stil = styles.add_style(
                "AbschnittTitel", WD_STYLE_TYPE.PARAGRAPH
            )
        except ValueError:
            abschnitt_stil = styles["AbschnittTitel"]
        abschnitt_stil.font.size = Pt(14)
        abschnitt_stil.font.bold = True
        abschnitt_stil.font.color.rgb = RGBColor(42, 42, 42)
        abschnitt_stil.paragraph_format.space_before = Pt(20)
        abschnitt_stil.paragraph_format.space_after = Pt(12)
        abschnitt_stil.paragraph_format.keep_with_next = True

        # Fließtext-Stil
        try:
            fliesstext_stil = styles.add_style(
                "Fliesstext", WD_STYLE_TYPE.PARAGRAPH
            )
        except ValueError:
            fliesstext_stil = styles["Fliesstext"]
        schriftgroesse = ebook.stilvorlage.schriftgroesse_text
        fliesstext_stil.font.size = Pt(schriftgroesse)
        fliesstext_stil.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fliesstext_stil.paragraph_format.space_after = Pt(8)
        fliesstext_stil.paragraph_format.line_spacing_rule = (
            WD_LINE_SPACING.ONE_POINT_FIVE
        )
        fliesstext_stil.paragraph_format.first_line_indent = Cm(0.5)

        # Fließtext ohne Einzug
        try:
            fliesstext_erster = styles.add_style(
                "FliesstextErster", WD_STYLE_TYPE.PARAGRAPH
            )
        except ValueError:
            fliesstext_erster = styles["FliesstextErster"]
        fliesstext_erster.font.size = Pt(schriftgroesse)
        fliesstext_erster.paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
        )
        fliesstext_erster.paragraph_format.space_after = Pt(8)
        fliesstext_erster.paragraph_format.line_spacing_rule = (
            WD_LINE_SPACING.ONE_POINT_FIVE
        )

        # Widmung-Stil
        try:
            widmung_stil = styles.add_style("Widmung", WD_STYLE_TYPE.PARAGRAPH)
        except ValueError:
            widmung_stil = styles["Widmung"]
        widmung_stil.font.size = Pt(12)
        widmung_stil.font.italic = True
        widmung_stil.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        widmung_stil.paragraph_format.space_before = Pt(100)
        widmung_stil.paragraph_format.space_after = Pt(100)

        # Bildunterschrift-Stil
        try:
            bildunter_stil = styles.add_style(
                "Bildunterschrift", WD_STYLE_TYPE.PARAGRAPH
            )
        except ValueError:
            bildunter_stil = styles["Bildunterschrift"]
        bildunter_stil.font.size = Pt(9)
        bildunter_stil.font.italic = True
        bildunter_stil.font.color.rgb = RGBColor(74, 74, 74)
        bildunter_stil.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bildunter_stil.paragraph_format.space_before = Pt(6)
        bildunter_stil.paragraph_format.space_after = Pt(12)

        # Copyright-Stil
        try:
            copyright_stil = styles.add_style(
                "Copyright", WD_STYLE_TYPE.PARAGRAPH
            )
        except ValueError:
            copyright_stil = styles["Copyright"]
        copyright_stil.font.size = Pt(9)
        copyright_stil.font.color.rgb = RGBColor(106, 106, 106)
        copyright_stil.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # IVZ-Überschrift
        try:
            ivz_titel = styles.add_style("IVZTitel", WD_STYLE_TYPE.PARAGRAPH)
        except ValueError:
            ivz_titel = styles["IVZTitel"]
        ivz_titel.font.size = Pt(20)
        ivz_titel.font.bold = True
        ivz_titel.paragraph_format.space_after = Pt(20)

        # IVZ-Eintrag
        try:
            ivz_eintrag = styles.add_style(
                "IVZEintrag", WD_STYLE_TYPE.PARAGRAPH
            )
        except ValueError:
            ivz_eintrag = styles["IVZEintrag"]
        ivz_eintrag.font.size = Pt(11)
        ivz_eintrag.paragraph_format.space_before = Pt(4)
        ivz_eintrag.paragraph_format.space_after = Pt(4)

    def _seitenformat_einrichten(self, ebook: EBook) -> None:
        """Richtet das Seitenformat ein."""
        section = self.document.sections[0]

        # Seitengröße
        groesse = self.SEITENFORMATE.get(
            ebook.stilvorlage.seitenformat, (14.8, 21.0)
        )
        section.page_width = Cm(groesse[0])
        section.page_height = Cm(groesse[1])

        # Seitenränder
        raender = ebook.stilvorlage.seitenraender
        section.top_margin = Cm(raender[0])
        section.right_margin = Cm(raender[1])
        section.bottom_margin = Cm(raender[2])
        section.left_margin = Cm(raender[3])

    def _eigenschaften_setzen(self, ebook: EBook) -> None:
        """Setzt die Dokumenteigenschaften."""
        props = self.document.core_properties
        props.title = ebook.metadaten.titel
        props.author = ebook.metadaten.autor
        props.subject = ebook.metadaten.beschreibung
        props.keywords = ", ".join(ebook.metadaten.schluesselwoerter)
        props.language = ebook.metadaten.sprache

    def _seitenumbruch(self) -> None:
        """Fügt einen Seitenumbruch ein."""
        self.document.add_page_break()

    def _titelseite_erstellen(self, ebook: EBook) -> None:
        """Erstellt die Titelseite."""
        # Abstand oben
        for _ in range(3):
            self.document.add_paragraph()

        # Coverbild falls vorhanden
        if ebook.coverbild and ebook.coverbild.exists():
            try:
                para = self.document.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run()
                run.add_picture(str(ebook.coverbild), width=Cm(8))
                self.document.add_paragraph()
            except Exception:
                pass  # Bild konnte nicht geladen werden

        # Titel
        self.document.add_paragraph(
            ebook.metadaten.titel, style="BuchTitel"
        )

        # Untertitel
        if ebook.metadaten.untertitel:
            self.document.add_paragraph(
                ebook.metadaten.untertitel, style="BuchUntertitel"
            )

        # Abstand
        for _ in range(2):
            self.document.add_paragraph()

        # Autor
        self.document.add_paragraph(ebook.metadaten.autor, style="BuchAutor")

        # Verlag
        if ebook.metadaten.verlag:
            for _ in range(3):
                self.document.add_paragraph()
            self.document.add_paragraph(
                ebook.metadaten.verlag, style="Copyright"
            )

    def _copyright_seite_erstellen(self, ebook: EBook) -> None:
        """Erstellt die Copyright-Seite."""
        # Abstand oben
        for _ in range(20):
            self.document.add_paragraph()

        # Copyright
        self.document.add_paragraph(
            ebook.metadaten.copyright, style="Copyright"
        )

        # ISBN
        if ebook.metadaten.isbn:
            self.document.add_paragraph()
            self.document.add_paragraph(
                f"ISBN: {ebook.metadaten.isbn}", style="Copyright"
            )

        # Erscheinungsdatum
        self.document.add_paragraph()
        datum = ebook.metadaten.erscheinungsdatum.strftime("%B %Y")
        self.document.add_paragraph(f"Erschienen: {datum}", style="Copyright")

    def _widmung_erstellen(self, ebook: EBook) -> None:
        """Erstellt die Widmungsseite."""
        self.document.add_paragraph(ebook.widmung, style="Widmung")

    def _inhaltsverzeichnis_erstellen(self, ebook: EBook) -> None:
        """Erstellt das Inhaltsverzeichnis."""
        self.document.add_paragraph("Inhaltsverzeichnis", style="IVZTitel")

        for kapitel in ebook.kapitel:
            # Kapiteleintrag
            para = self.document.add_paragraph(style="IVZEintrag")
            run = para.add_run(f"Kapitel {kapitel.kapitelnummer}: ")
            run.bold = True
            para.add_run(kapitel.titel)

            # Abschnitte eingerückt
            for abschnitt in kapitel.abschnitte:
                para = self.document.add_paragraph(style="IVZEintrag")
                para.paragraph_format.left_indent = Cm(1)
                para.add_run(f"• {abschnitt.ueberschrift}")

    def _vorwort_erstellen(self, ebook: EBook) -> None:
        """Erstellt das Vorwort."""
        self.document.add_paragraph("Vorwort", style="KapitelTitel")

        absaetze = ebook.vorwort.strip().split("\n\n")
        for i, absatz in enumerate(absaetze):
            if absatz.strip():
                stil = "FliesstextErster" if i == 0 else "Fliesstext"
                self._formatierter_absatz(absatz.strip(), stil)

    def _nachwort_erstellen(self, ebook: EBook) -> None:
        """Erstellt das Nachwort."""
        self.document.add_paragraph("Nachwort", style="KapitelTitel")

        absaetze = ebook.nachwort.strip().split("\n\n")
        for i, absatz in enumerate(absaetze):
            if absatz.strip():
                stil = "FliesstextErster" if i == 0 else "Fliesstext"
                self._formatierter_absatz(absatz.strip(), stil)

    def _kapitel_erstellen(self, kapitel) -> None:
        """Erstellt ein Kapitel."""
        # Kapitelüberschrift
        titel_text = f"Kapitel {kapitel.kapitelnummer}: {kapitel.titel}"
        self.document.add_paragraph(titel_text, style="KapitelTitel")

        # Kapitelinhalt
        if kapitel.inhalt:
            absaetze = kapitel.inhalt.strip().split("\n\n")
            for i, absatz in enumerate(absaetze):
                if absatz.strip():
                    stil = "FliesstextErster" if i == 0 else "Fliesstext"
                    self._formatierter_absatz(absatz.strip(), stil)

        # Abschnitte
        for abschnitt in kapitel.abschnitte:
            self._abschnitt_erstellen(abschnitt)

    def _abschnitt_erstellen(self, abschnitt) -> None:
        """Erstellt einen Abschnitt."""
        # Abschnittsüberschrift
        self.document.add_paragraph(
            abschnitt.ueberschrift, style="AbschnittTitel"
        )

        # Abschnittsinhalt
        if abschnitt.inhalt:
            absaetze = abschnitt.inhalt.strip().split("\n\n")
            for i, absatz in enumerate(absaetze):
                if absatz.strip():
                    stil = "FliesstextErster" if i == 0 else "Fliesstext"
                    self._formatierter_absatz(absatz.strip(), stil)

        # Bilder
        for bild in abschnitt.bilder:
            self._bild_erstellen(bild)

    def _bild_erstellen(self, bild) -> None:
        """Erstellt ein Bild mit optionaler Bildunterschrift."""
        if not bild.pfad.exists():
            return

        try:
            para = self.document.add_paragraph()

            # Ausrichtung
            if bild.ausrichtung == "links":
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif bild.ausrichtung == "rechts":
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = para.add_run()

            # Bildgröße
            if bild.breite:
                breite = Pt(bild.breite)
            else:
                breite = Cm(10)

            if bild.hoehe:
                run.add_picture(str(bild.pfad), width=breite, height=Pt(bild.hoehe))
            else:
                run.add_picture(str(bild.pfad), width=breite)

            # Bildunterschrift
            if bild.bildunterschrift:
                self.document.add_paragraph(
                    bild.bildunterschrift, style="Bildunterschrift"
                )

        except Exception:
            pass  # Bild konnte nicht verarbeitet werden

    def _formatierter_absatz(self, text: str, stil: str) -> None:
        """
        Erstellt einen formatierten Absatz mit Markdown-Unterstützung.

        Unterstützt:
        - **fett**
        - *kursiv*
        - __unterstrichen__
        """
        para = self.document.add_paragraph(style=stil)

        # Muster für Formatierung
        muster = r"(\*\*.*?\*\*|\*.*?\*|__.*?__)"
        teile = re.split(muster, text)

        for teil in teile:
            if not teil:
                continue

            if teil.startswith("**") and teil.endswith("**"):
                # Fett
                run = para.add_run(teil[2:-2])
                run.bold = True
            elif teil.startswith("*") and teil.endswith("*"):
                # Kursiv
                run = para.add_run(teil[1:-1])
                run.italic = True
            elif teil.startswith("__") and teil.endswith("__"):
                # Unterstrichen
                run = para.add_run(teil[2:-2])
                run.underline = True
            else:
                para.add_run(teil)
